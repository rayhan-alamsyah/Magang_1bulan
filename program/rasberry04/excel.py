# ==========================================================
# EXCEL.PY
# DATABASE QC MENGGUNAKAN EXCEL
# ==========================================================

from pathlib import Path
from uuid import uuid4

from openpyxl import (
    Workbook,
    load_workbook
)

from openpyxl.styles import (
    Font,
    Alignment
)

from openpyxl.drawing.image import (
    Image as ExcelImage
)


# ==========================================================
# PATH
# ==========================================================

BASE_DIR = Path(__file__).resolve().parent
ROOT_DIR = BASE_DIR.parent.parent

DATA_DIR = ROOT_DIR / "data"
EXCEL_FILE = DATA_DIR / "QC.xlsx"
PHOTO_DIR = DATA_DIR / "foto"


# ==========================================================
# BUAT FOLDER
# ==========================================================

def buat_folder():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    PHOTO_DIR.mkdir(parents=True, exist_ok=True)


# ==========================================================
# BUAT EXCEL
# ==========================================================

def buat_excel():
    buat_folder()

    if EXCEL_FILE.exists():
        return

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "QC"

    sheet["A1"] = "REKAP"
    sheet["H1"] = "DATA"

    # ------------------------------------------------------
    # HEADER REKAP
    # ------------------------------------------------------

    rekap_headers = [
        "Tanggal",
        "GOOD",
        "NG Upper Kiri",
        "NG Upper Kanan",
        "NG Lower Kiri",
        "NG Lower Kanan"
    ]

    for column, value in enumerate(rekap_headers, 1):
        cell = sheet.cell(2, column)
        cell.value = value

        cell.font = Font(bold=True)

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # ------------------------------------------------------
    # HEADER DATA
    # ------------------------------------------------------

    data_headers = [
        "ID",
        "No",
        "Tanggal",
        "Waktu",
        "Foto",
        "Tinggi Kiri",
        "Tinggi Kanan",
        "Upper Kiri",
        "Lower Kiri",
        "Upper Kanan",
        "Lower Kanan",
        "Hasil"
    ]

    for column, value in enumerate(data_headers, 8):
        cell = sheet.cell(2, column)
        cell.value = value

        cell.font = Font(bold=True)

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # ------------------------------------------------------
    # LEBAR KOLOM
    # ------------------------------------------------------

    widths = {
        "A": 15,
        "B": 10,
        "C": 18,
        "D": 18,
        "E": 18,
        "F": 18,
        "G": 3,
        "H": 10,
        "I": 10,
        "J": 15,
        "K": 12,
        "L": 30,
        "M": 15,
        "N": 15,
        "O": 15,
        "P": 15,
        "Q": 15,
        "R": 15,
        "S": 22
    }

    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    # Sembunyikan kolom ID
    sheet.column_dimensions["H"].hidden = True

    workbook.save(EXCEL_FILE)
    workbook.close()


# ==========================================================
# BACA DATA
# ==========================================================

def baca_data():
    buat_excel()

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["QC"]

    data = []

    for row in range(3, sheet.max_row + 1):

        if sheet.cell(row, 8).value is None:
            continue

        data.append({
            "id": sheet.cell(row, 8).value,
            "no": sheet.cell(row, 9).value,
            "tanggal": sheet.cell(row, 10).value,
            "waktu": sheet.cell(row, 11).value,
            "foto": sheet.cell(row, 12).value,
            "tinggi_kiri": sheet.cell(row, 13).value,
            "tinggi_kanan": sheet.cell(row, 14).value,
            "upper_kiri": sheet.cell(row, 15).value,
            "lower_kiri": sheet.cell(row, 16).value,
            "upper_kanan": sheet.cell(row, 17).value,
            "lower_kanan": sheet.cell(row, 18).value,
            "hasil": sheet.cell(row, 19).value
        })

    workbook.close()

    return data


# ==========================================================
# HITUNG REKAP
# ==========================================================

def hitung_rekap():

    rekap = {}

    for item in baca_data():

        tanggal = str(
            item.get("tanggal") or "-"
        )

        if tanggal not in rekap:

            rekap[tanggal] = {
                "good": 0,
                "upper_kiri": 0,
                "upper_kanan": 0,
                "lower_kiri": 0,
                "lower_kanan": 0
            }

        hasil = str(
            item.get("hasil") or ""
        ).upper()

        if hasil == "GOOD":
            rekap[tanggal]["good"] += 1

        for key in [
            "upper_kiri",
            "upper_kanan",
            "lower_kiri",
            "lower_kanan"
        ]:

            status = str(
                item.get(key) or ""
            ).upper()

            if status == "NG":
                rekap[tanggal][key] += 1

    return rekap


# ==========================================================
# UPDATE REKAP
# ==========================================================

def update_rekap():

    buat_excel()

    rekap = hitung_rekap()

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["QC"]

    # ------------------------------------------------------
    # HAPUS REKAP LAMA
    # ------------------------------------------------------

    for row in range(
        3,
        sheet.max_row + 1
    ):

        for column in range(1, 7):
            sheet.cell(
                row,
                column
            ).value = None

    # ------------------------------------------------------
    # TULIS REKAP BARU
    # ------------------------------------------------------

    row = 3

    for tanggal, nilai in sorted(
        rekap.items(),
        reverse=True
    ):

        values = [
            tanggal,
            nilai["good"],
            nilai["upper_kiri"],
            nilai["upper_kanan"],
            nilai["lower_kiri"],
            nilai["lower_kanan"]
        ]

        for column, value in enumerate(
            values,
            1
        ):

            sheet.cell(
                row,
                column
            ).value = value

        row += 1

    workbook.save(EXCEL_FILE)
    workbook.close()


# ==========================================================
# CEK FOTO
# ==========================================================

def _safe_photo(relative):

    if not relative or str(relative) == "-":
        return None

    root = PHOTO_DIR.resolve()

    path = (
        PHOTO_DIR / str(relative)
    ).resolve()

    if (
        path.is_file()
        and root in path.parents
    ):
        return path

    return None


# ==========================================================
# EMBED FOTO KE EXCEL
# ==========================================================

def _embed_photo(
    sheet,
    row,
    relative
):

    path = _safe_photo(relative)

    if path is None:
        return

    try:

        image = ExcelImage(str(path))

        image.width = 180
        image.height = 120

        sheet.add_image(
            image,
            f"L{row}"
        )

        sheet.row_dimensions[row].height = 95

    except Exception as error:

        print(
            "Gagal embed foto Excel:",
            error
        )


# ==========================================================
# SIMPAN DATA
# ==========================================================

def simpan_data(data):

    buat_excel()

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["QC"]

    data_id = str(uuid4())

    # ------------------------------------------------------
    # NOMOR
    # ------------------------------------------------------

    nomor = 1

    for row in range(
        3,
        sheet.max_row + 1
    ):

        value = sheet.cell(
            row,
            9
        ).value

        if isinstance(value, int):

            nomor = max(
                nomor,
                value + 1
            )

    # ------------------------------------------------------
    # BARIS BARU
    # ------------------------------------------------------

    row = max(
        3,
        sheet.max_row + 1
    )

    values = [
        data_id,
        nomor,
        data.get("tanggal"),
        data.get("waktu"),
        data.get("foto", "-"),
        data.get("tinggi_kiri"),
        data.get("tinggi_kanan"),
        data.get("upper_kiri"),
        data.get("lower_kiri"),
        data.get("upper_kanan"),
        data.get("lower_kanan"),
        data.get("hasil")
    ]

    for column, value in enumerate(
        values,
        8
    ):

        sheet.cell(
            row,
            column
        ).value = value

    # ------------------------------------------------------
    # EMBED FOTO
    # ------------------------------------------------------

    _embed_photo(
        sheet,
        row,
        data.get("foto")
    )

    # ------------------------------------------------------
    # ALIGNMENT
    # ------------------------------------------------------

    for column in range(
        8,
        20
    ):

        sheet.cell(
            row,
            column
        ).alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    workbook.save(EXCEL_FILE)
    workbook.close()

    update_rekap()

    return data_id


# ==========================================================
# EDIT DATA
# ==========================================================

def edit_data(data):

    buat_excel()

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["QC"]

    target_row = None

    for row in range(
        3,
        sheet.max_row + 1
    ):

        value = sheet.cell(
            row,
            8
        ).value

        if str(value) == str(
            data.get("id")
        ):

            target_row = row
            break

    if target_row is None:

        workbook.close()

        return False

    fields = [
        (10, "tanggal"),
        (11, "waktu"),
        (13, "tinggi_kiri"),
        (14, "tinggi_kanan"),
        (15, "upper_kiri"),
        (16, "lower_kiri"),
        (17, "upper_kanan"),
        (18, "lower_kanan"),
        (19, "hasil")
    ]

    for column, key in fields:

        sheet.cell(
            target_row,
            column
        ).value = data.get(key)

    workbook.save(EXCEL_FILE)
    workbook.close()

    update_rekap()

    return True


# ==========================================================
# HAPUS SATU DATA
# ==========================================================

def hapus_data(data_id):

    buat_excel()

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["QC"]

    target_row = None
    photo = None

    # ------------------------------------------------------
    # CARI DATA BERDASARKAN ID
    # ------------------------------------------------------

    for row in range(
        3,
        sheet.max_row + 1
    ):

        value = sheet.cell(
            row,
            8
        ).value

        if str(value) == str(data_id):

            target_row = row

            photo = sheet.cell(
                row,
                12
            ).value

            break

    if target_row is None:

        workbook.close()

        return False

    # ======================================================
    # SIMPAN INFORMASI GAMBAR SEBELUM DELETE ROW
    # ======================================================

    gambar = []

    for image in list(sheet._images):

        try:

            image_row = (
                image.anchor._from.row + 1
            )

            if image_row == target_row:

                gambar.append({
                    "image": image,
                    "row": image_row
                })

        except Exception as error:

            print(
                "Gagal membaca posisi gambar:",
                error
            )

    # ======================================================
    # HAPUS BARIS DATA
    # ======================================================

    sheet.delete_rows(
        target_row,
        1
    )

    # ======================================================
    # HAPUS GAMBAR MILIK DATA TERSEBUT
    # ======================================================

    for item in gambar:

        image = item["image"]

        try:

            if image in sheet._images:
                sheet._images.remove(image)

        except Exception as error:

            print(
                "Gagal menghapus gambar Excel:",
                error
            )

    # ======================================================
    # GESER GAMBAR YANG ADA DI BAWAHNYA
    # ======================================================

    for image in list(sheet._images):

        try:

            old_row = (
                image.anchor._from.row + 1
            )

            if old_row >= target_row:

                new_row = old_row - 1

                image.anchor._from.row = (
                    new_row - 1
                )

                # Kalau gambar menggunakan TwoCellAnchor
                if hasattr(
                    image.anchor,
                    "_to"
                ):

                    image.anchor._to.row = (
                        image.anchor._to.row - 1
                    )

        except Exception as error:

            print(
                "Gagal menggeser gambar:",
                error
            )

    # ======================================================
    # RAPIAKAN TINGGI BARIS
    # ======================================================

    for row in range(
        3,
        sheet.max_row + 1
    ):

        sheet.row_dimensions[row].height = 15

    # Baris yang masih mempunyai gambar dibuat tinggi
    for image in list(sheet._images):

        try:

            image_row = (
                image.anchor._from.row + 1
            )

            sheet.row_dimensions[
                image_row
            ].height = 95

        except Exception:
            pass

    # ======================================================
    # SIMPAN EXCEL
    # ======================================================

    workbook.save(EXCEL_FILE)
    workbook.close()

    # ======================================================
    # HAPUS FILE FOTO DARI FOLDER
    # ======================================================

    photo_path = _safe_photo(photo)

    if photo_path:

        try:

            photo_path.unlink()

        except Exception as error:

            print(
                "Gagal menghapus foto:",
                error
            )

    # ======================================================
    # UPDATE REKAP
    # ======================================================

    update_rekap()

    return True


# ==========================================================
# HAPUS SEMUA DATA BERDASARKAN TANGGAL
# ==========================================================

def hapus_tanggal(tanggal):

    buat_excel()

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["QC"]

    rows = []
    photos = []

    # ------------------------------------------------------
    # CARI SEMUA BARIS TANGGAL
    # ------------------------------------------------------

    for row in range(
        3,
        sheet.max_row + 1
    ):

        tanggal_excel = sheet.cell(
            row,
            10
        ).value

        if str(tanggal_excel) == str(tanggal):

            rows.append(row)

            photos.append(
                sheet.cell(
                    row,
                    12
                ).value
            )

    if not rows:

        workbook.close()

        return False

    # ======================================================
    # SIMPAN GAMBAR SEBELUM DELETE ROW
    # ======================================================

    gambar_hapus = []

    for image in list(sheet._images):

        try:

            image_row = (
                image.anchor._from.row + 1
            )

            if image_row in rows:

                gambar_hapus.append({
                    "image": image,
                    "row": image_row
                })

        except Exception as error:

            print(
                "Gagal membaca posisi gambar:",
                error
            )

    # ======================================================
    # HAPUS BARIS DARI BAWAH KE ATAS
    # ======================================================

    for row in reversed(rows):

        sheet.delete_rows(
            row,
            1
        )

    # ======================================================
    # HAPUS GAMBAR YANG BERHUBUNGAN
    # ======================================================

    for item in gambar_hapus:

        image = item["image"]

        try:

            if image in sheet._images:

                sheet._images.remove(image)

        except Exception as error:

            print(
                "Gagal menghapus gambar Excel:",
                error
            )

    # ======================================================
    # GESER GAMBAR YANG TERSISA
    # ======================================================

    for image in list(sheet._images):

        try:

            old_row = (
                image.anchor._from.row + 1
            )

            # Hitung berapa baris yang dihapus
            # berada di atas gambar
            jumlah_hapus_di_atas = 0

            for deleted_row in rows:

                if deleted_row < old_row:
                    jumlah_hapus_di_atas += 1

            if jumlah_hapus_di_atas > 0:

                new_row = (
                    old_row -
                    jumlah_hapus_di_atas
                )

                image.anchor._from.row = (
                    new_row - 1
                )

                # Kalau menggunakan TwoCellAnchor
                if hasattr(
                    image.anchor,
                    "_to"
                ):

                    image.anchor._to.row = (
                        image.anchor._to.row -
                        jumlah_hapus_di_atas
                    )

        except Exception as error:

            print(
                "Gagal menggeser gambar:",
                error
            )

    # ======================================================
    # RAPIAKAN TINGGI BARIS
    # ======================================================

    for row in range(
        3,
        sheet.max_row + 1
    ):

        sheet.row_dimensions[row].height = 15

    # Baris yang masih memiliki gambar
    # dibuat tinggi kembali
    for image in list(sheet._images):

        try:

            image_row = (
                image.anchor._from.row + 1
            )

            sheet.row_dimensions[
                image_row
            ].height = 95

        except Exception:
            pass

    # ======================================================
    # SIMPAN EXCEL
    # ======================================================

    workbook.save(EXCEL_FILE)
    workbook.close()

    # ======================================================
    # HAPUS FILE FOTO
    # ======================================================

    for photo in photos:

        photo_path = _safe_photo(photo)

        if photo_path:

            try:

                photo_path.unlink()

            except Exception as error:

                print(
                    "Gagal menghapus foto:",
                    error
                )

    # ======================================================
    # HAPUS FOLDER TANGGAL JIKA KOSONG
    # ======================================================

    folder = PHOTO_DIR / str(tanggal)

    try:

        if (
            folder.is_dir()
            and not any(folder.iterdir())
        ):

            folder.rmdir()

    except Exception:

        pass

    # ======================================================
    # UPDATE REKAP
    # ======================================================

    update_rekap()

    return True


# ==========================================================
# LOKASI EXCEL
# ==========================================================

def lokasi_excel():

    return str(EXCEL_FILE)